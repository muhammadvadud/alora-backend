from __future__ import annotations

import csv
import io
from contextvars import ContextVar
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO

from django.db import transaction

from apps.catalog.models import Category, Product, ProductVariant

IMPORT_IN_PROGRESS: ContextVar[bool] = ContextVar("catalog_import_in_progress", default=False)

MAX_ROWS = 5000
MAX_FILE_SIZE = 5 * 1024 * 1024
REQUIRED_COLUMNS = {
    "product_code",
    "product_name",
    "category_slug",
    "variant_sku",
    "color_name",
    "price",
}
_LEGACY_COLUMNS = {
    "code": "product_code",
    "name": "product_name",
    "sku": "variant_sku",
}


@dataclass(frozen=True)
class ImportReport:
    created: int
    updated: int
    skipped: int
    errors: list[dict[str, Any]]


def _error(row: int, field: str, message: str) -> dict[str, Any]:
    return {"row": row, "field": field, "message": message}


def _normalise_row(row: dict[str, Any]) -> dict[str, str]:
    result = {str(key).strip(): str(value or "").strip() for key, value in row.items() if key}
    for old, new in _LEGACY_COLUMNS.items():
        if new not in result and old in result:
            result[new] = result[old]
    return result


def _read_rows(file: BinaryIO) -> tuple[list[dict[str, Any]], list[str]]:
    name = str(getattr(file, "name", "")).lower()
    raw = file.read()
    if len(raw) > MAX_FILE_SIZE:
        raise ValueError("Fayl hajmi 5 MB dan oshmasligi kerak")
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency is installed in production
            raise ValueError("XLSX importi uchun openpyxl o'rnatilishi kerak") from exc
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            raise ValueError("XLSX faylda faol varaq topilmadi")
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        return [dict(zip(headers, values, strict=False)) for values in rows], headers
    if name and "." in name and not name.endswith(".csv"):
        raise ValueError("Faqat CSV yoki XLSX fayl qabul qilinadi")
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return list(reader), list(reader.fieldnames or [])


def _validate_headers(headers: list[str]) -> None:
    normalised = set(headers) | {_LEGACY_COLUMNS.get(header, header) for header in headers}
    missing = REQUIRED_COLUMNS - normalised
    if missing:
        raise ValueError(f"Kerakli ustunlar topilmadi: {', '.join(sorted(missing))}")


def _decimal(value: str, field: str) -> Decimal:
    try:
        result = Decimal(value.replace(",", "."))
    except (InvalidOperation, AttributeError):
        raise ValueError(f"{field} son bo'lishi kerak") from None
    if result < 0:
        raise ValueError(f"{field} manfiy bo'lmasligi kerak")
    return result


def _bool(value: str) -> bool:
    return value.lower() not in {"0", "false", "no", "yo'q", "yoq", "inactive"}


def _upsert_row(*, tenant, row: dict[str, str]) -> tuple[bool, bool]:
    for field in REQUIRED_COLUMNS:
        if not row.get(field):
            raise ValueError(f"{field} kiritilishi shart")
    price = _decimal(row["price"], "price")
    category, _ = Category.objects.get_or_create(
        tenant=tenant,
        slug=row["category_slug"],
        defaults={"name": row["category_slug"]},
    )
    product, product_created = Product.objects.update_or_create(
        tenant=tenant,
        code=row["product_code"],
        defaults={
            "category": category,
            "name": row["product_name"],
            "style": row.get("style", ""),
            "pattern": row.get("pattern", ""),
            "segment": row.get("segment", ""),
            "recommended_rooms": [
                x.strip() for x in row.get("recommended_rooms", "").split(";") if x.strip()
            ],
            "is_active": _bool(row.get("is_active", "true")),
        },
    )
    _variant, variant_created = ProductVariant.objects.update_or_create(
        tenant=tenant,
        sku=row["variant_sku"],
        defaults={
            "product": product,
            "color_name": row["color_name"],
            "price": price,
            "price_unit": row.get("price_unit") or "per_sqm",
            "width_m": _decimal(row["width_m"], "width_m") if row.get("width_m") else None,
            "is_active": _bool(row.get("is_active", "true")),
        },
    )
    return product_created, variant_created


def import_products(file: BinaryIO, tenant, dry_run: bool = False) -> ImportReport:
    """CSV/XLSX fayldagi mahsulot va variantlarni bitta tranzaksiyada import qiladi."""
    rows, headers = _read_rows(file)
    _validate_headers(headers)
    if len(rows) > MAX_ROWS:
        raise ValueError("Faylda 5000 qatordan ko'p bo'lmasligi kerak")
    created = updated = skipped = 0
    errors: list[dict[str, Any]] = []
    marker = IMPORT_IN_PROGRESS.set(True)
    try:
        with transaction.atomic():
            for row_number, raw_row in enumerate(rows, start=2):
                try:
                    product_created, variant_created = _upsert_row(
                        tenant=tenant, row=_normalise_row(raw_row)
                    )
                    if product_created or variant_created:
                        created += 1
                    else:
                        updated += 1
                except (ValueError, InvalidOperation) as exc:
                    skipped += 1
                    message = str(exc)
                    field = message.split(" ", 1)[0] if message else "row"
                    errors.append(_error(row_number, field, message))
            if dry_run:
                transaction.set_rollback(True)
        report = ImportReport(created=created, updated=updated, skipped=skipped, errors=errors)
    finally:
        IMPORT_IN_PROGRESS.reset(marker)
    if not dry_run and (created or updated):
        from apps.chat.grpc_server.catalog_service import publish_catalog_delta

        publish_catalog_delta(
            tenant_id=str(tenant.id), entity="catalog", entity_id=str(tenant.id), action="update"
        )
    return report


def import_catalog_csv(*, tenant, file, dry_run: bool = False) -> ImportReport:
    """Eski servis nomi uchun moslik wrapper'i."""
    return import_products(file=file, tenant=tenant, dry_run=dry_run)
